"""Build synthetic drawing PDFs for the test suite.

Real client drawings cannot live in the repository, so the suite generates its
own. The sheets here deliberately reproduce the awkward parts of the real
fixture set:

* a media box that is **not** at the origin, like the Lami Architects set
  whose box runs (-1192, -842) to (1192, 842)
* title block values that sit *below* their label (`Drawing No.`) as well as
  *to the right* of it (`Scale`)
* multi-page files, encrypted files, corrupt files and image-only sheets

Anything that needs to be true of a real drawing should be reproduced here
first, so the behaviour is pinned by a test rather than by one lucky file.
"""

from __future__ import annotations

import io
import zlib
from dataclasses import dataclass, field
from pathlib import Path

import pikepdf

# A1 landscape in points, the size the real fixtures use.
A1_WIDTH = 2384.0
A1_HEIGHT = 1684.0
A3_WIDTH = 1191.0
A3_HEIGHT = 842.0


@dataclass(slots=True)
class SheetSpec:
    """One drawing sheet: what its title block says."""

    drawing_no: str = "A-101"
    title: str = "GROUND FLOOR PLAN"
    revision: str = "C"
    scale: str = "1 : 100"
    project_no: str = "LM2426"
    #: Extra text placed in the drawing area, away from the title block.
    body: list[str] = field(default_factory=lambda: ["GRID A", "GRID B", "3000", "3200"])
    #: Omit the title block entirely, to exercise the unidentified path.
    include_title_block: bool = True
    width: float = A1_WIDTH
    height: float = A1_HEIGHT


def _escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _title_block_items(spec: SheetSpec) -> list[tuple[float, float, float, str]]:
    """Title block text as (x, y, font size, text), in page-local coordinates.

    Laid out like the real fixture: bottom-right corner, labels above their
    values, except `Scale` whose value sits to the right of the label.
    """
    right = spec.width - 60
    items: list[tuple[float, float, float, str]] = [
        # Company block, top right
        (right - 250, spec.height - 60, 11, "LAMI ARCHITECTS"),
        (right - 250, spec.height - 160, 8, "L E A D  A R C H I T E C T"),
        # Project and title, above the number block
        (right - 250, 300, 8, "P R O J E C T"),
        (right - 250, 280, 11, "Al Basateen Farm - Villa"),
        (right - 250, 200, 8, "Drawing Title"),
        (right - 250, 175, 12, spec.title),
        # Scale: value to the RIGHT of the label
        (right - 250, 120, 8, "Scale"),
        (right - 180, 120, 9, spec.scale),
        # Number block: values BELOW their labels
        (right - 250, 90, 8, "Project No."),
        (right - 250, 72, 10, spec.project_no),
        (right - 110, 90, 8, "Drawing No."),
        (right - 110, 72, 12, spec.drawing_no),
        (right - 20, 90, 8, "Rev."),
        (right - 20, 72, 12, spec.revision),
    ]
    return items


def _content_stream(spec: SheetSpec, x0: float, y0: float) -> bytes:
    parts: list[str] = []

    items: list[tuple[float, float, float, str]] = []
    if spec.include_title_block:
        items.extend(_title_block_items(spec))

    # Body text, spread across the left two thirds of the sheet.
    for index, text in enumerate(spec.body):
        items.append((spec.width * 0.15, spec.height * 0.7 - index * 40, 14, text))

    for local_x, local_y, size, text in items:
        parts.append(
            f"BT /F1 {size} Tf {x0 + local_x:.2f} {y0 + local_y:.2f} Td ({_escape(text)}) Tj ET"
        )

    # A border, so the sheet is not a blank page.
    parts.append(
        f"0.5 w {x0 + 20:.2f} {y0 + 20:.2f} {spec.width - 40:.2f} {spec.height - 40:.2f} re S"
    )
    return "\n".join(parts).encode("latin-1")


def build_pdf(
    path: str | Path,
    sheets: list[SheetSpec] | None = None,
    *,
    offset_origin: bool = True,
    password: str | None = None,
    producer: str = "C-Quest test fixture",
) -> Path:
    """Write a PDF containing one page per sheet spec.

    `offset_origin` centres the media box on (0, 0), reproducing the real
    fixture. Leave it on unless a test specifically needs a page at origin.
    """
    sheets = sheets or [SheetSpec()]
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    pdf = pikepdf.Pdf.new()
    font = pdf.make_indirect(
        pikepdf.Dictionary(
            Type=pikepdf.Name.Font,
            Subtype=pikepdf.Name.Type1,
            BaseFont=pikepdf.Name.Helvetica,
            Encoding=pikepdf.Name.WinAnsiEncoding,
        )
    )

    for spec in sheets:
        x0 = -spec.width / 2 if offset_origin else 0.0
        y0 = -spec.height / 2 if offset_origin else 0.0
        stream = pdf.make_stream(_content_stream(spec, x0, y0))
        page = pikepdf.Dictionary(
            Type=pikepdf.Name.Page,
            MediaBox=[x0, y0, x0 + spec.width, y0 + spec.height],
            Resources=pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font)),
            Contents=stream,
        )
        pdf.pages.append(pikepdf.Page(pdf.make_indirect(page)))

    with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
        meta["pdf:Producer"] = producer

    if password is None:
        pdf.save(target)
    else:
        pdf.save(target, encryption=pikepdf.Encryption(owner=password, user=password, R=6))
    pdf.close()
    return target


def build_scanned_pdf(path: str | Path, pages: int = 1) -> Path:
    """A page holding one large image and no text, like a photocopied sheet."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    pdf = pikepdf.Pdf.new()
    width, height = A3_WIDTH, A3_HEIGHT
    pixels, rows = 200, 140
    raw = bytes([220]) * (pixels * rows)  # flat grey, compresses to nothing

    for _ in range(pages):
        image = pdf.make_stream(zlib.compress(raw))
        image.Type = pikepdf.Name.XObject
        image.Subtype = pikepdf.Name.Image
        image.Width = pixels
        image.Height = rows
        image.ColorSpace = pikepdf.Name.DeviceGray
        image.BitsPerComponent = 8
        image.Filter = pikepdf.Name.FlateDecode

        content = pdf.make_stream(
            f"q {width:.2f} 0 0 {height:.2f} 0 0 cm /Im0 Do Q".encode("latin-1")
        )
        page = pikepdf.Dictionary(
            Type=pikepdf.Name.Page,
            MediaBox=[0, 0, width, height],
            Resources=pikepdf.Dictionary(XObject=pikepdf.Dictionary(Im0=image)),
            Contents=content,
        )
        pdf.pages.append(pikepdf.Page(pdf.make_indirect(page)))

    pdf.save(target)
    pdf.close()
    return target


def build_corrupt_pdf(path: str | Path) -> Path:
    """A file that starts like a PDF and then stops. Must land in quarantine."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    buffer = io.BytesIO()
    build_pdf_to_buffer(buffer)
    data = buffer.getvalue()
    target.write_bytes(data[: len(data) // 3])  # truncated: no xref, no trailer
    return target


def build_pdf_to_buffer(buffer: io.BytesIO, sheets: list[SheetSpec] | None = None) -> None:
    """Same as :func:`build_pdf` but into memory, used by the corrupt builder."""
    sheets = sheets or [SheetSpec()]
    pdf = pikepdf.Pdf.new()
    font = pdf.make_indirect(
        pikepdf.Dictionary(
            Type=pikepdf.Name.Font,
            Subtype=pikepdf.Name.Type1,
            BaseFont=pikepdf.Name.Helvetica,
            Encoding=pikepdf.Name.WinAnsiEncoding,
        )
    )
    for spec in sheets:
        stream = pdf.make_stream(_content_stream(spec, 0.0, 0.0))
        page = pikepdf.Dictionary(
            Type=pikepdf.Name.Page,
            MediaBox=[0, 0, spec.width, spec.height],
            Resources=pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font)),
            Contents=stream,
        )
        pdf.pages.append(pikepdf.Page(pdf.make_indirect(page)))
    pdf.save(buffer)
    pdf.close()


# ── Whole fixture sets ─────────────────────────────────────────────────

#: A small architectural set, the shape of `01_normal`.
NORMAL_SET: list[tuple[str, str, str]] = [
    ("A-001", "SITE PLAN", "C"),
    ("A-101", "GROUND FLOOR PLAN", "C"),
    ("A-102", "FIRST FLOOR PLAN", "C"),
    ("A-103", "ROOF PLAN", "C"),
    ("A-104", "GROUND FLOOR REFLECTED CEILING PLAN", "C"),
    ("A-105", "FIRST FLOOR REFLECTED CEILING PLAN", "C"),
    ("A-201", "NORTH AND SOUTH ELEVATIONS", "C"),
    ("A-202", "EAST AND WEST ELEVATIONS", "C"),
    ("A-301", "SECTION A-A", "C"),
    ("A-302", "SECTION B-B", "C"),
]


def build_normal_pair(root: str | Path) -> tuple[Path, Path]:
    """Two full issues: three drawings revised, one new, everything else the same.

    Returns the (old, new) folders.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"

    for number, title, revision in NORMAL_SET:
        build_pdf(
            old_dir / f"{number}-{title.title().replace(' ', '-')}-Rev{revision}.pdf",
            [SheetSpec(drawing_no=number, title=title, revision=revision)],
        )

    for number, title, revision in NORMAL_SET:
        # A-101, A-102 and A-201 move to revision D; the rest stay at C.
        new_revision = "D" if number in {"A-101", "A-102", "A-201"} else revision
        body = ["GRID A", "GRID B", "3200", "3400"] if new_revision == "D" else None
        spec = SheetSpec(drawing_no=number, title=title, revision=new_revision)
        if body:
            spec.body = body
        build_pdf(
            new_dir / f"{number}-{title.title().replace(' ', '-')}-Rev{new_revision}.pdf",
            [spec],
        )

    build_pdf(
        new_dir / "A-303-Section-C-C-RevA.pdf",
        [SheetSpec(drawing_no="A-303", title="SECTION C-C", revision="A")],
    )

    return old_dir, new_dir


def build_partial_issue_pair(root: str | Path, old_count: int = 40) -> tuple[Path, Path]:
    """A partial issue: a large previous set, and only three sheets reissued.

    This is the case that decides whether people trust the tool. The new set
    must never be reported as hundreds of removals.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"

    for index in range(old_count):
        number = f"A-{100 + index}"
        build_pdf(
            old_dir / f"{number}-Rev-C.pdf",
            [SheetSpec(drawing_no=number, title=f"PLAN {index}", revision="C")],
        )

    for index in range(3):
        number = f"A-{100 + index}"
        build_pdf(
            new_dir / f"{number}-Rev-D.pdf",
            [SheetSpec(drawing_no=number, title=f"PLAN {index}", revision="D")],
        )

    return old_dir, new_dir


def build_messy_folder(root: str | Path) -> Path:
    """Junk files, an archive subfolder, a duplicate, and a nested set."""
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)

    build_pdf(root / "A-101-Rev-C.pdf", [SheetSpec(drawing_no="A-101", revision="C")])
    build_pdf(
        root / "Architectural" / "A-102-Rev-C.pdf",
        [SheetSpec(drawing_no="A-102", title="FIRST FLOOR PLAN", revision="C")],
    )
    build_pdf(
        root / "Architectural" / "Levels" / "A-103-Rev-C.pdf",
        [SheetSpec(drawing_no="A-103", title="ROOF PLAN", revision="C")],
    )

    # A true duplicate: identical bytes in two places.
    duplicate_source = root / "A-101-Rev-C.pdf"
    (root / "Copy of A-101-Rev-C.pdf").write_bytes(duplicate_source.read_bytes())

    # Superseded material that must be skipped whole.
    build_pdf(
        root / "superseded" / "A-101-Rev-B.pdf", [SheetSpec(drawing_no="A-101", revision="B")]
    )
    build_pdf(root / "_archive" / "A-999-Rev-A.pdf", [SheetSpec(drawing_no="A-999")])

    # Junk the scanner must exclude.
    (root / "Thumbs.db").write_bytes(b"not a drawing")
    (root / "desktop.ini").write_text("[.ShellClassInfo]", encoding="utf-8")
    (root / "~$register.xlsx").write_bytes(b"office lock file")

    # Other files worth reporting but not scanning as drawings.
    (root / "A-101.dwg").write_bytes(b"DWG placeholder")
    (root / "drawing-list.xlsx").write_bytes(b"XLSX placeholder")

    return root


# ── Drawing lists ──────────────────────────────────────────────────────


def build_clean_drawing_list(path: str | Path) -> Path:
    """A tidy register: header on row 1, one row per drawing."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Register"
    sheet.append(["Drawing No", "Title", "Rev"])
    for number, title, revision in NORMAL_SET:
        sheet.append([number, title, revision])

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target


def build_messy_drawing_list(path: str | Path) -> Path:
    """The register as a human actually formats it.

    Reproduces what real files do: a logo block above the header, the header
    on row 7, a decoy worksheet, discipline headings inside the data, and
    blank separator rows.
    """
    import openpyxl

    book = openpyxl.Workbook()

    # A decoy sheet, of the kind every real workbook has.
    notes = book.active
    notes.title = "Notes"
    notes.append(["Project information"])
    notes.append(["Issued by", "Lami Architects"])
    notes.append(["Contact", "info@example.com"])

    sheet = book.create_sheet("Drawing Register")
    sheet.append(["ACME ARCHITECTS"])  # 1: logo row
    sheet.append([])  # 2
    sheet.append(["Project:", "Al Basateen Farm - Villa"])  # 3
    sheet.append(["Issue:", "IFC Rev D"])  # 4
    sheet.append(["Date:", "27/06/2025"])  # 5
    sheet.append([])  # 6
    sheet.append(["Dwg No.", "Sheet Name", "Rev.", "Status"])  # 7: the header
    sheet.append(["ARCHITECTURAL"])  # a discipline heading
    for number, title, revision in NORMAL_SET[:5]:
        sheet.append([number, title, revision, "For construction"])
    sheet.append([])  # separator
    sheet.append(["STRUCTURAL"])
    for number, title, revision in NORMAL_SET[5:]:
        sheet.append([number, title, revision, "For construction"])

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target


def build_revision_matrix_list(path: str | Path) -> Path:
    """A register with no Rev column: one column per issue date instead."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Register"
    sheet.append(["Drawing No", "Title", "12/01/2025", "27/06/2025", "01/09/2026"])
    sheet.append(["A-101", "GROUND FLOOR PLAN", "A", "B", "C"])
    sheet.append(["A-102", "FIRST FLOOR PLAN", "A", "B", ""])
    sheet.append(["A-103", "ROOF PLAN", "A", "", ""])

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target


# ── Phase 3 fixtures: `03_renamed`, `08_naming_mess`, `09_collision` ────


#: The ten drawings of `03_renamed` shared by both issues, as
#: (old number, new number, title, body lines). The body of each drawing uses
#: words no other drawing shares, so content fingerprints of two different
#: drawings stay well below the 0.4 line while the old and new versions of one
#: drawing — identical body, identical title, identical sheet size — match at
#: 1.0. The body lines read as room names and notes pinned to each sheet.
RENAMED_SET: list[tuple[str, str, str, list[str]]] = [
    (
        "UVU-ARC-001",
        "UVU-KEO-XX-03-DR-A-0001",
        "GROUND FLOOR PLAN",
        ["TYPICAL ROOM LAYOUT", "WALL PARTITION HEIGHT 3001"],
    ),
    (
        "UVU-ARC-002",
        "UVU-KEO-XX-03-DR-A-0002",
        "ROOF PLAN",
        ["ROOF SLOPE FALL OUTLET", "PARAPET GUTTER 3002"],
    ),
    (
        "UVU-ARC-003",
        "UVU-KEO-XX-03-DR-A-0003",
        "EXTERNAL WALL SECTION",
        ["CAVITY INSULATION CLADDING", "TIE VAPOUR BARRIER 3003"],
    ),
    (
        "UVU-ARC-004",
        "UVU-KEO-XX-03-DR-A-0004",
        "STAIRCASE DETAILS",
        ["STAIR FLIGHT LANDING", "HANDRAIL RISER TREAD 3004"],
    ),
    (
        "UVU-ARC-005",
        "UVU-KEO-XX-03-DR-A-0005",
        "TOILET LAYOUT",
        ["TOILET SHOWER COMPARTMENT", "VANITY PIPEWORK VENT 3005"],
    ),
    (
        "UVU-ARC-006",
        "UVU-KEO-XX-03-DR-A-0006",
        "KITCHEN PLAN",
        ["KITCHEN BENCH HOOD", "EXTRACT DUCT SERVICE 3006"],
    ),
    (
        "UVU-ARC-007",
        "UVU-KEO-XX-03-DR-A-0007",
        "ENTRANCE LOBBY",
        ["ENTRANCE LOBBY RECEPTION", "DESK GLAZING SCREEN 3007"],
    ),
    (
        "UVU-ARC-008",
        "UVU-KEO-XX-03-DR-A-0008",
        "FOUNDATION PLAN",
        ["FOUNDATION STRIP CONCRETE", "REINFORCEMENT BLINDING BEAM 3008"],
    ),
    (
        "UVU-ARC-009",
        "UVU-KEO-XX-03-DR-A-0009",
        "FIRE ESCAPE ROUTES",
        ["CORRIDOR FIRE ESCAPE", "ROUTE EXIT LAMP 3009"],
    ),
    (
        "UVU-ARC-010",
        "UVU-KEO-XX-03-DR-A-0010",
        "PLANT DECK LAYOUT",
        ["PLANT DECK UNIT", "CONDENSER LOUVRE GRILLE 3010"],
    ),
]

#: A drawing present only in the old issue of `03_renamed`, so the matcher's
#: old-unmatched list has something real to hold.
RENAMED_OLD_ONLY: tuple[str, str, list[str]] = (
    "UVU-ARC-011",
    "OFFICE FURNITURE PLAN",
    ["FURNITURE SCHEDULE ITEM", "FINISH ANNOTATION 3011"],
)

#: A drawing present only in the new issue of `03_renamed`, so the matcher's
#: new-unmatched list has something real to hold.
RENAMED_NEW_ONLY: tuple[str, str, list[str]] = (
    "UVU-KEO-XX-03-DR-A-0011",
    "SIGNAGE SCHEDULE",
    ["WAYFINDING ARROW MARKING", "NOTICE ZONE 3012"],
)


def _renamed_old_filename(drawing_no: str, title: str) -> str:
    """The old-standard file name: the title travels with the number."""
    return f"{drawing_no}_{title.replace(' ', '_')}_RevC.pdf"


def _renamed_new_filename(drawing_no: str) -> str:
    """The new-standard file name: number only, no title."""
    return f"{drawing_no}_RevC.pdf"


def build_renamed_pair(root: str | Path) -> tuple[Path, Path]:
    """`03_renamed`: the same ten drawings under a new naming standard.

    The client changed the standard mid-project. Both sides carry the same ten
    drawings with identical title, body text and sheet size; only the drawing
    number (and so the file name) changes, `UVU-ARC-00N` on the old side and
    `UVU-KEO-XX-03-DR-A-000N` on the new. One extra drawing exists only on the
    old side and one only on the new side, so unmatched lists are exercised.
    Returns the (old, new) folders.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"

    for old_no, new_no, title, body in RENAMED_SET:
        revision = "C"
        build_pdf(
            old_dir / _renamed_old_filename(old_no, title),
            [SheetSpec(drawing_no=old_no, title=title, body=body, revision=revision)],
        )
        build_pdf(
            new_dir / _renamed_new_filename(new_no),
            [SheetSpec(drawing_no=new_no, title=title, body=body, revision=revision)],
        )

    old_no, old_title, old_body = RENAMED_OLD_ONLY
    build_pdf(
        old_dir / _renamed_old_filename(old_no, old_title),
        [SheetSpec(drawing_no=old_no, title=old_title, body=old_body, revision="C")],
    )

    new_no, new_title, new_body = RENAMED_NEW_ONLY
    build_pdf(
        new_dir / _renamed_new_filename(new_no),
        [SheetSpec(drawing_no=new_no, title=new_title, body=new_body, revision="C")],
    )

    return old_dir, new_dir


def build_naming_mess(root: str | Path) -> Path:
    """`08_naming_mess`: one folder full of naming sins on the same drawing.

    Every PDF is the same drawing — drawing number UVU-ARC-001 — saved under a
    name a real issue folder would contain: Windows copy junk, separators and
    case, a revision and status stamp, a date, a trailing space-dot, and en
    dashes that look exactly like hyphens. AGGRESSIVE cleaning must collapse
    all of them to ``uvuarc001``. Returns the ``new`` folder.
    """
    root = Path(root)
    new_dir = root / "new"
    spec = SheetSpec(drawing_no="UVU-ARC-001", title="GROUND FLOOR PLAN", revision="C")

    names = [
        "Copy of UVU-ARC-001.pdf",
        "UVU ARC 001 (1).pdf",
        "UVU_ARC_001_RevD_FOR APPROVAL.pdf",
        "01 - UVU-ARC-001 - 12.04.2026.pdf",
        "uvu-arc-001 FINAL final.pdf",
        "UVU\u2013ARC\u2013001.pdf",  # en dashes, not hyphens
        "UVU-ARC-001 .pdf",  # trailing space-dot
    ]
    for name in names:
        build_pdf(new_dir / name, [spec])

    return new_dir


def build_collision_folder(root: str | Path) -> Path:
    """`09_collision`: two different drawings whose cleaned names collide.

    Both received files read as drawing ``UVU-ARC-001`` once cleaned — the
    sequence prefix on the first and the date stamp on the second disappear at
    MEDIUM — so a naive rename that takes the number from the file name sends
    both to the same target. The title blocks tell the truth: they carry
    different drawing numbers (``UVU-ARC-001`` vs ``UVU-ARC-002``) and
    disjoint body text, so they are clearly two different drawings. Returns
    the ``new`` folder.
    """
    root = Path(root)
    new_dir = root / "new"

    build_pdf(
        new_dir / "01 - UVU-ARC-001.pdf",
        [
            SheetSpec(
                drawing_no="UVU-ARC-001",
                title="GROUND FLOOR PLAN",
                body=["ROOM 101 LAYOUT", "DOOR SCHEDULE D1 D2"],
            )
        ],
    )
    build_pdf(
        new_dir / "UVU-ARC-001 - 12.04.2026.pdf",
        [
            SheetSpec(
                drawing_no="UVU-ARC-002",
                title="ROOF PLAN",
                body=["ROOF DRAINAGE DETAIL", "FALL 1 IN 60 OUTLET"],
            )
        ],
    )

    return new_dir


# ── Phase 4 fixtures: `10_alignment` ─────────────────────────────────

#: Content band of a rich drawing page, as fractions of the sheet
#: (x0, y0, x1, y1). The band clears the bottom-right title-block strip, so a
#: text-anchor matcher that excludes that strip still sees the real content.
ALIGN_CONTENT_REGION = (0.08, 0.2, 0.92, 0.95)

#: The bottom-right strip that holds the title block, as fractions of the
#: sheet. Anchor matching must exclude every text item whose origin lands in
#: it: that text is identical sheet furniture on every issue of every sheet.
ALIGN_TITLEBLOCK_ZONE = (0.7, 0.0, 1.0, 0.12)

#: How many unique labels a rich drawing page carries: RM-01 .. RM-24.
ALIGN_LABEL_COUNT = 24
#: The labels form a 6-column x 4-row matrix inside the content band.
ALIGN_LABEL_COLUMNS = 6
ALIGN_LABEL_ROWS = 4

#: Content band of the rescaled pair's old sheet, centred on the page centre.
#: Scaling it 2.0 about the centre lands it on (0.1, 0.1)-(0.9, 0.9), so the
#: re-issued drawing never leaves the same-size media box.
RESCALED_CONTENT_REGION = (0.3, 0.3, 0.7, 0.7)


def align_label_names(
    prefix: str = "RM", count: int = ALIGN_LABEL_COUNT, start: int = 1
) -> list[str]:
    """Label strings ``RM-01`` .. ``RM-24``, zero padded to two digits."""
    return [f"{prefix}-{n:02d}" for n in range(start, start + count)]


def align_label_fractions(
    index: int,
    region: tuple[float, float, float, float] = ALIGN_CONTENT_REGION,
    *,
    variant: int = 0,
) -> tuple[float, float]:
    """The baseline origin (fx, fy) of label *index*, as sheet fractions.

    Layout ``variant=0`` (the rich drawing pages) fills the content band with a
    6 x 4 matrix, columns fastest: RM-01 sits bottom-left of the band, RM-06
    bottom-right, RM-24 top-right. Label text is drawn with its baseline
    origin exactly on the returned point, so downstream anchor code can
    recover every label's expected position from :data:`ALIGN_LABEL_COLUMNS`,
    :data:`ALIGN_LABEL_ROWS` and the band.

    Layout ``variant=1`` (the S-101 alternative) uses a 4 x 4 matrix with the
    same conventions, so its labels never share coordinates with layout 0.
    """
    rx0, ry0, rx1, ry1 = region
    if variant == 0:
        columns, rows = ALIGN_LABEL_COLUMNS, ALIGN_LABEL_ROWS
        col = index % columns
        row = index // columns
    else:
        columns = rows = 4
        col = index % columns
        row = index // columns
    fx = rx0 + (col + 0.5) * (rx1 - rx0) / columns
    fy = ry0 + (row + 0.5) * (ry1 - ry0) / rows
    return fx, fy


@dataclass(slots=True)
class DrawingPageSpec:
    """One rich, CAD-like page of the Phase 4 alignment fixtures.

    Everything the drawing needs is explicit: its identity (drawing number,
    revision, scale), the labels that will become text anchors, and the
    geometry knobs that make one fixture case differ from another.
    """

    drawing_no: str = "A-101"
    revision: str = "C"
    scale: str = "1 : 100"
    labels: tuple[str, ...] = field(default_factory=lambda: tuple(align_label_names()))
    #: Extra annotation lines, stacked down the left edge of the content band.
    content_lines: tuple[str, ...] = ()
    #: Draw the fine grid/hatch lines. Off gives a bubble-free detail sheet.
    grid: bool = True
    circles: bool = True
    #: 0 = the standard layout, 1 = the S-101 alternative (see the geometry
    #: tables below: different grid treatment, rectangles and hatch zones).
    variant: int = 0
    #: The content band, as fractions of the sheet.
    region: tuple[float, float, float, float] = ALIGN_CONTENT_REGION
    label_size: float = 9.0
    width: float = A1_WIDTH
    height: float = A1_HEIGHT


def drawing_page_spec(
    content_lines: list[str] | None = None,
    grid: bool = True,
    **overrides: object,
) -> DrawingPageSpec:
    """Build a :class:`DrawingPageSpec` with the common knobs as keywords.

    ``content_lines`` become annotation lines stacked down the left of the
    content band; ``grid=False`` removes every fine line, hatch and circle.
    Any other :class:`DrawingPageSpec` field may be passed as a keyword.
    """
    values: dict[str, object] = {
        "content_lines": tuple(content_lines) if content_lines else (),
        "grid": grid,
    }
    values.update(overrides)
    return DrawingPageSpec(**values)


def _op_text(x: float, y: float, size: float, text: str) -> str:
    return f"BT /F1 {size:.1f} Tf {x:.2f} {y:.2f} Td ({_escape(text)}) Tj ET"


def _op_line(x0: float, y0: float, x1: float, y1: float, width: float) -> str:
    return f"{width:.2f} w {x0:.2f} {y0:.2f} m {x1:.2f} {y1:.2f} l S"


def _op_rect(x: float, y: float, w: float, h: float, width: float) -> str:
    return f"{width:.2f} w {x:.2f} {y:.2f} {w:.2f} {h:.2f} re S"


def _op_circle(cx: float, cy: float, r: float, width: float) -> str:
    """A circle as four cubic segments (no PDF arc operator needed)."""
    k = 0.5522847498
    ops = [f"{width:.2f} w {cx + r:.2f} {cy:.2f} m"]
    for x1, y1, x2, y2, x3, y3 in (
        (cx + r, cy + k * r, cx + k * r, cy + r, cx, cy + r),
        (cx - k * r, cy + r, cx - r, cy + k * r, cx - r, cy),
        (cx - r, cy - k * r, cx - k * r, cy - r, cx, cy - r),
        (cx + k * r, cy - r, cx + r, cy - k * r, cx + r, cy),
    ):
        ops.append(f"{x1:.2f} {y1:.2f} {x2:.2f} {y2:.2f} {x3:.2f} {y3:.2f} c")
    ops.append("S")
    return "\n".join(ops)


#: Region-local rectangles (u0, v0, u1, v1) of the standard layout.
_VARIANT0_RECTS = (
    (0.06, 0.06, 0.30, 0.24),
    (0.44, 0.10, 0.62, 0.30),
    (0.70, 0.06, 0.94, 0.20),
    (0.08, 0.44, 0.24, 0.62),
    (0.36, 0.40, 0.56, 0.58),
    (0.64, 0.42, 0.92, 0.58),
    (0.10, 0.72, 0.28, 0.92),
    (0.50, 0.70, 0.68, 0.90),
    (0.76, 0.68, 0.94, 0.88),
)

#: Region-local rectangles of the S-101 alternative; (u0, v0, u1, v1, filled).
_VARIANT1_RECTS = (
    (0.08, 0.08, 0.46, 0.30, False),
    (0.60, 0.10, 0.94, 0.34, False),
    (0.16, 0.10, 0.34, 0.28, True),  # a solid block: nowhere in the old sheet
    (0.10, 0.50, 0.30, 0.72, False),
    (0.44, 0.46, 0.64, 0.90, False),
    (0.78, 0.56, 0.94, 0.74, False),
    (0.44, 0.66, 0.60, 0.80, True),
)

#: Squares (region-local) that carry 45-degree hatching on the standard sheet.
_VARIANT0_HATCH = ((0.055, 0.055, 0.295, 0.245), (0.77, 0.69, 0.935, 0.875))

#: "Grid bubble" circles (region-local centres) on the standard sheet.
_VARIANT0_CIRCLES = ((0.18, 0.94), (0.50, 0.94), (0.82, 0.94))


def _hatch_ops(x0: float, y0: float, x1: float, y1: float, slope: int = 1) -> list[str]:
    """Fine 45-degree lines filling the rectangle, slope +1 or -1."""
    dx = x1 - x0
    dy = y1 - y0
    step = min(dx, dy) / 24.0
    lines: list[str] = []
    runs = int((dx + dy) / step) + 3
    for m in range(runs):
        start_x = x0 - dy + m * step
        if start_x + dy <= x0 or start_x >= x1:
            continue
        sx = max(start_x, x0)
        ex = min(start_x + dy, x1)
        if slope == 1:
            lines.append(_op_line(sx, y0 + (sx - start_x), ex, y0 + (ex - start_x), 0.18))
        else:
            lines.append(_op_line(sx, y1 - (sx - start_x), ex, y1 - (ex - start_x), 0.18))
    return lines


def _content_ops(spec: DrawingPageSpec) -> list[str]:
    """The drawing proper: content band, fine lines, shapes and labels."""
    w, h = spec.width, spec.height
    rx0, ry0, rx1, ry1 = spec.region
    x0, y0 = rx0 * w, ry0 * h
    x1, y1 = rx1 * w, ry1 * h
    cw, ch = x1 - x0, y1 - y0
    ops: list[str] = [_op_rect(x0, y0, cw, ch, 0.5)]

    if spec.grid:
        if spec.variant == 0:
            for i in range(1, 32):  # verticals every 1/32 of the band width
                ops.append(_op_line(x0 + cw * i / 32, y0, x0 + cw * i / 32, y1, 0.18))
            for j in range(1, 16):  # horizontals every 1/16 of the band height
                ops.append(_op_line(x0, y0 + ch * j / 16, x1, y0 + ch * j / 16, 0.18))
            for hx0, hy0, hx1, hy1 in _VARIANT0_HATCH:
                ops.extend(_hatch_ops(x0 + hx0 * cw, y0 + hy0 * ch, x0 + hx1 * cw, y0 + hy1 * ch))
            for hx0, hy0, hx1, hy1 in _VARIANT0_HATCH:
                ops.extend(
                    _hatch_ops(x0 + hx0 * cw, y0 + hy0 * ch, x0 + hx1 * cw, y0 + hy1 * ch, slope=-1)
                )
        else:
            # The S-101 sheet: a dense 45-degree lattice over the whole band.
            ops.extend(_hatch_ops(x0, y0, x1, y1))
            ops.extend(_hatch_ops(x0, y0, x1, y1, slope=-1))

    rects = _VARIANT1_RECTS if spec.variant == 1 else _VARIANT0_RECTS
    for rect in rects:
        if len(rect) == 5:
            u0, v0, u1, v1, filled = rect
        else:
            u0, v0, u1, v1 = rect
            filled = False
        rx, ry = x0 + u0 * cw, y0 + v0 * ch
        rw, rh = (u1 - u0) * cw, (v1 - v0) * ch
        if filled:
            ops.append(f"{rx:.2f} {ry:.2f} {rw:.2f} {rh:.2f} re f")
        else:
            ops.append(_op_rect(rx, ry, rw, rh, 0.6))

    if spec.circles and spec.variant == 0:
        radius = 0.0125 * w
        for u, v in _VARIANT0_CIRCLES:
            ops.append(_op_circle(x0 + u * cw, y0 + v * ch, radius, 0.5))

    for index, label in enumerate(spec.labels):
        fx, fy = align_label_fractions(index, spec.region, variant=spec.variant)
        ops.append(_op_text(fx * w, fy * h, spec.label_size, label))

    for index, line in enumerate(spec.content_lines):
        fy = y1 - (index + 1) * 0.03 * h
        if fy < y0 + 0.02 * h:
            break
        ops.append(_op_text(x0 + 8, fy, 8.0, line))
    return ops


def _sheet_furniture_ops(
    width: float, height: float, drawing_no: str, revision: str, scale: str
) -> list[str]:
    """The sheet frame and the minimal bottom-right title block, in points.

    Laid out like the real fixture: ``Drawing No.`` and ``Rev.`` labels with
    their values *below* them, ``Scale`` with its value to the right. The
    block is white-filled first so it reads cleanly when drawn over other
    content (the rescaled re-issue does exactly that).
    """
    x0 = 0.70 * width
    x1 = width - 20.0
    y0 = 20.0
    y1 = 0.12 * height
    ops: list[str] = [f"0.8 w 20.00 20.00 {width - 40:.2f} {height - 40:.2f} re S"]

    # White sheet behind the block, then the block border.
    ops.append(f"1 g {x0:.2f} {y0:.2f} {x1 - x0:.2f} {y1 - y0:.2f} re f 0 g")
    ops.append(f"0.6 w {x0:.2f} {y0:.2f} {x1 - x0:.2f} {y1 - y0:.2f} re S")

    divider_drawing = 0.845 * width
    divider_rev = 0.955 * width
    row_divider = 0.058 * height
    ops.append(f"0.3 w {divider_drawing:.2f} {y0:.2f} {divider_drawing:.2f} {y1:.2f} S")
    ops.append(f"0.3 w {divider_rev:.2f} {y0:.2f} {divider_rev:.2f} {y1:.2f} S")
    ops.append(f"0.3 w {x0:.2f} {row_divider:.2f} {x1:.2f} {row_divider:.2f} S")

    text_left = x0 + 0.01 * width
    label_y = 0.096 * height
    value_y = 0.078 * height
    scale_y = 0.038 * height
    rev_centre = (divider_drawing + divider_rev) / 2

    ops.append(_op_text(text_left, label_y, 7.5, "Drawing No."))
    ops.append(_op_text(rev_centre - 10, label_y, 7.5, "Rev."))
    ops.append(_op_text(text_left, value_y, 11.0, drawing_no))
    ops.append(_op_text(rev_centre - 4, value_y, 11.0, revision))
    ops.append(_op_text(text_left, scale_y, 7.5, "Scale"))
    ops.append(_op_text(text_left + 150, scale_y, 11.0, scale))
    return ops


def _drawing_ops(spec: DrawingPageSpec) -> list[str]:
    ops = list(
        _sheet_furniture_ops(spec.width, spec.height, spec.drawing_no, spec.revision, spec.scale)
    )
    ops.extend(_content_ops(spec))
    return ops


def _make_font(pdf: pikepdf.Pdf) -> pikepdf.Object:
    return pdf.make_indirect(
        pikepdf.Dictionary(
            Type=pikepdf.Name.Font,
            Subtype=pikepdf.Name.Type1,
            BaseFont=pikepdf.Name.Helvetica,
            Encoding=pikepdf.Name.WinAnsiEncoding,
        )
    )


def write_drawing_pdf(
    path: str | Path,
    *,
    drawing_no: str = "A-101",
    revision: str = "C",
    scale: str = "1 : 100",
    labels: list[str] | None = None,
    content_lines: list[str] | None = None,
    grid: bool = True,
    circles: bool = True,
    variant: int = 0,
    region: tuple[float, float, float, float] = ALIGN_CONTENT_REGION,
    label_size: float = 9.0,
    width: float = A1_WIDTH,
    height: float = A1_HEIGHT,
) -> Path:
    """Write one rich drawing page and return its path.

    The labels default to ``RM-01`` .. ``RM-24`` and are drawn with Helvetica
    (base-14, WinAnsi, ASCII only), so pdfium extracts every label as text.
    Like the rest of this module the page media box is centred on the origin;
    all coordinates inside are sheet-local fractions. Deterministic for equal
    arguments.
    """
    spec = DrawingPageSpec(
        drawing_no=drawing_no,
        revision=revision,
        scale=scale,
        labels=tuple(labels) if labels is not None else tuple(align_label_names()),
        content_lines=tuple(content_lines) if content_lines else (),
        grid=grid,
        circles=circles,
        variant=variant,
        region=region,
        label_size=label_size,
        width=width,
        height=height,
    )
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    pdf = pikepdf.Pdf.new()
    font = _make_font(pdf)
    x0 = -width / 2
    y0 = -height / 2
    body = "\n".join(_drawing_ops(spec)).encode("latin-1")
    stream = pdf.make_stream(f"q 1 0 0 1 {x0:.2f} {y0:.2f} cm\n".encode("latin-1") + body + b"\nQ")
    page = pikepdf.Dictionary(
        Type=pikepdf.Name.Page,
        MediaBox=[x0, y0, x0 + width, y0 + height],
        Resources=pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font)),
        Contents=stream,
    )
    pdf.pages.append(pikepdf.Page(pdf.make_indirect(page)))

    with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
        meta["pdf:Producer"] = "C-Quest test fixture"
    pdf.save(target)
    pdf.close()
    return target


def _num(value: float) -> str:
    """A PDF number literal: no exponent, up to four decimals, no trailing zeros."""
    text = f"{value:.4f}".rstrip("0").rstrip(".")
    return text or "0"


def wrap_content(
    pdf_path: str | Path,
    matrix6: tuple[float, float, float, float, float, float],
    out_path: str | Path,
    *,
    resize_media: tuple[float, float] | None = None,
) -> Path:
    """Copy *pdf_path* to *out_path*, wrapping each page's content stream as
    ``q <matrix6> cm <original ops> Q``.

    ``matrix6`` is a PDF cm matrix ``(a, b, c, d, e, f)`` in user-space
    points. The fixture pages use a media box centred on the origin, so the
    page centre *is* the user-space origin: a pure translation shifts the
    content along the sheet, and a rotation about the origin is a rotation
    about the page centre. ``resize_media=(w, h)`` replaces the media box with
    one of that size centred on the old box's centre. The output is
    deterministic for equal inputs. Returns *out_path*.
    """
    src = Path(pdf_path)
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    a, b, c, d, e, f = (_num(v) for v in matrix6)
    prefix = f"q {a} {b} {c} {d} {e} {f} cm\n".encode("latin-1")

    with pikepdf.open(src) as pdf:
        for page in pdf.pages:
            contents = page.obj.Contents
            if isinstance(contents, pikepdf.Array):
                parts = [stream.read_bytes() for stream in contents]
            else:
                parts = [contents.read_bytes()]
            page.obj.Contents = pdf.make_stream(prefix + b"\n".join(parts) + b"\nQ")
            if resize_media is not None:
                new_w, new_h = resize_media
                box = [float(v) for v in page.obj.MediaBox]
                cx = (box[0] + box[2]) / 2
                cy = (box[1] + box[3]) / 2
                page.obj.MediaBox = pikepdf.Array(
                    [cx - new_w / 2, cy - new_h / 2, cx + new_w / 2, cy + new_h / 2]
                )
        pdf.save(target)
    return target


def _append_content_ops(pdf_path: str | Path, out_path: str | Path, ops: list[str]) -> Path:
    """Save *pdf_path* to *out_path* with extra operators after its content.

    Used to put fresh sheet furniture (frame + title block) on top of a
    wrapped page whose own furniture scaled or rotated off the sheet. The
    output may be the input file itself (in-place append).
    """
    target = Path(out_path)
    body = "\n".join(ops).encode("latin-1")
    with pikepdf.open(Path(pdf_path), allow_overwriting_input=target == Path(pdf_path)) as pdf:
        page = pdf.pages[0]
        # The page's own content is written in sheet-local coordinates and
        # translated by the media box's lower-left corner; the overlay ops
        # (also sheet-local) need the same translation to land in the box.
        box = [float(v) for v in page.obj.MediaBox]
        overlay = f"q 1 0 0 1 {box[0]:.2f} {box[1]:.2f} cm\n".encode("latin-1") + body + b"\nQ"
        contents = page.obj.Contents
        if isinstance(contents, pikepdf.Array):
            existing = b"\n".join(stream.read_bytes() for stream in contents)
        else:
            existing = contents.read_bytes()
        page.obj.Contents = pdf.make_stream(existing + b"\n" + overlay)
        target.parent.mkdir(parents=True, exist_ok=True)
        pdf.save(target)
    return target


def _image_only_pdf(
    path: str | Path,
    gray_pixels: bytes,
    width_px: int,
    height_px: int,
    *,
    width_pt: float,
    height_pt: float,
) -> Path:
    """A full-page image PDF with no text layer, like a photocopied sheet."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    pdf = pikepdf.Pdf.new()
    image = pdf.make_stream(zlib.compress(gray_pixels))
    image.Type = pikepdf.Name.XObject
    image.Subtype = pikepdf.Name.Image
    image.Width = width_px
    image.Height = height_px
    image.ColorSpace = pikepdf.Name.DeviceGray
    image.BitsPerComponent = 8
    image.Filter = pikepdf.Name.FlateDecode

    content = pdf.make_stream(
        f"q {width_pt:.2f} 0 0 {height_pt:.2f} 0 0 cm /Im0 Do Q".encode("latin-1")
    )
    page = pikepdf.Dictionary(
        Type=pikepdf.Name.Page,
        MediaBox=[0, 0, width_pt, height_pt],
        Resources=pikepdf.Dictionary(XObject=pikepdf.Dictionary(Im0=image)),
        Contents=content,
    )
    pdf.pages.append(pikepdf.Page(pdf.make_indirect(page)))
    pdf.save(target)
    pdf.close()
    return target


def build_alignment_clean(root: str | Path) -> Path:
    """`10_alignment/clean_pair`: identical sheets except one changed label.

    The old sheet is A-101 Rev C with RM-01..RM-24. The new sheet is the same
    drawing re-issued as Rev D with exactly one real change: label RM-07 is
    now RM-70. Every other label, line and value keeps its position, so the
    pair needs the identity transform. Returns *root* with old/ and new/.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"
    write_drawing_pdf(old_dir / "A-101-RevC.pdf", drawing_no="A-101", revision="C")

    labels = align_label_names()
    labels[6] = "RM-70"  # index 6 is RM-07
    write_drawing_pdf(new_dir / "A-101-RevD.pdf", drawing_no="A-101", revision="D", labels=labels)
    return root


def build_alignment_shifted(root: str | Path) -> Path:
    """`10_alignment/shifted`: the same sheet, its content plotted 40 mm right.

    The whole old content stream is wrapped in a cm translation of
    +40 mm of paper (= 40 / 25.4 * 72 points) along the sheet's x axis; the
    media box does not change. Returns *root* with old/ and new/.
    """
    root = Path(root)
    shift_x = 40.0 / 25.4 * 72.0
    old_pdf = write_drawing_pdf(root / "old" / "A-101-RevC.pdf", drawing_no="A-101", revision="C")
    wrap_content(
        old_pdf, (1.0, 0.0, 0.0, 1.0, shift_x, 0.0), root / "new" / "A-101-RevC-shifted.pdf"
    )
    return root


def build_alignment_rescaled(root: str | Path) -> Path:
    """`10_alignment/rescaled`: A-101 at 1:100 re-issued with content doubled.

    The old sheet draws its content inside :data:`RESCALED_CONTENT_REGION`
    (centred on the page); the new sheet is the same content scaled 2.0 about
    the page centre on the same-size media box, re-labelled 1 : 50. Fresh
    sheet furniture (frame + title block) is drawn on top because the old
    furniture scaled off the sheet with the content. Returns *root*.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"
    write_drawing_pdf(
        old_dir / "A-101-RevC-1-100.pdf",
        drawing_no="A-101",
        revision="C",
        scale="1 : 100",
        region=RESCALED_CONTENT_REGION,
    )
    wrapped = wrap_content(
        old_dir / "A-101-RevC-1-100.pdf",
        (2.0, 0.0, 0.0, 2.0, 0.0, 0.0),
        new_dir / "A-101-RevD-1-50.pdf",
    )
    furniture = _sheet_furniture_ops(A1_WIDTH, A1_HEIGHT, "A-101", "D", "1 : 50")
    _append_content_ops(wrapped, new_dir / "A-101-RevD-1-50.pdf", furniture)
    return root


def build_alignment_rotated(root: str | Path) -> Path:
    """`10_alignment/rotated`: the same sheet, its content rotated 90 degrees.

    The whole old content stream is wrapped in a 90-degree rotation about the
    page centre (the cm matrix ``(0 1 -1 0 0 0)`` in the centred media box);
    the media box does not change, so content near the sheet corners clips,
    as it would for a genuinely rotated re-issue. Returns *root*.
    """
    root = Path(root)
    old_pdf = write_drawing_pdf(root / "old" / "A-101-RevC.pdf", drawing_no="A-101", revision="C")
    wrap_content(
        old_pdf, (0.0, 1.0, -1.0, 0.0, 0.0, 0.0), root / "new" / "A-101-RevC-rotated90.pdf"
    )
    return root


def build_alignment_page_rotated(root: str | Path) -> Path:
    """`10_alignment/page_rotated`: identical content, different /Rotate flag.

    The new file is a byte-for-byte copy of the old drawing with the page's
    ``/Rotate`` entry set to 90; every content operator is unchanged, so the
    two renders come out equal once rotation-normalised. Returns *root*.
    """
    root = Path(root)
    old_pdf = write_drawing_pdf(root / "old" / "A-101-RevC.pdf", drawing_no="A-101", revision="C")
    new_pdf = root / "new" / "A-101-RevC-rotateflag90.pdf"
    new_pdf.parent.mkdir(parents=True, exist_ok=True)
    new_pdf.write_bytes(old_pdf.read_bytes())
    with pikepdf.open(new_pdf, allow_overwriting_input=True) as pdf:
        pdf.pages[0].Rotate = 90
        pdf.save(new_pdf)
    return root


def build_alignment_scanned(root: str | Path) -> Path:
    """`10_alignment/scanned`: the old sheet printed out and scanned back in.

    The old side is a normal A3 drawing page. The new side is that page
    rendered at ~150 dpi by pdfium, rotated 0.8 degrees with cv2 (white
    borders left by the rotation) and embedded full-page into a new PDF that
    has no text layer at all.
    """
    root = Path(root)
    old_pdf = write_drawing_pdf(
        root / "old" / "A-101-RevC.pdf",
        drawing_no="A-101",
        revision="C",
        width=A3_WIDTH,
        height=A3_HEIGHT,
    )

    import cv2

    from engine.utils.pdf_runtime import open_document

    with open_document(old_pdf) as document:
        bitmap = document[0].render(scale=150.0 / 72.0)
        pixels = bitmap.to_numpy()
    if pixels.ndim == 2:
        gray = pixels
    elif pixels.shape[2] == 3:
        gray = cv2.cvtColor(pixels, cv2.COLOR_BGR2GRAY)
    else:
        gray = cv2.cvtColor(pixels, cv2.COLOR_BGRA2GRAY)

    rows, cols = gray.shape
    matrix = cv2.getRotationMatrix2D((cols / 2, rows / 2), 0.8, 1.0)
    skewed = cv2.warpAffine(
        gray,
        matrix,
        (cols, rows),
        flags=cv2.INTER_LINEAR,
        borderMode=cv2.BORDER_CONSTANT,
        borderValue=255,
    )
    _image_only_pdf(
        root / "new" / "A-101-RevC-scanned.pdf",
        skewed.tobytes(),
        cols,
        rows,
        width_pt=A3_WIDTH,
        height_pt=A3_HEIGHT,
    )
    return root


def build_alignment_no_grid(root: str | Path) -> Path:
    """`10_alignment/no_grid`: a detail sheet with no grid, hatch or circles.

    Both sides are the identical drawing: labels and rectangles only, nothing
    a grid-bubble detector could mistake for a bubble. Returns *root*.
    """
    root = Path(root)
    pdf = write_drawing_pdf(
        root / "old" / "A-101-RevC.pdf",
        drawing_no="A-101",
        revision="C",
        grid=False,
        circles=False,
    )
    new_pdf = root / "new" / "A-101-RevC.pdf"
    new_pdf.parent.mkdir(parents=True, exist_ok=True)
    new_pdf.write_bytes(pdf.read_bytes())
    return root


def build_alignment_sparse_text(root: str | Path) -> Path:
    """`10_alignment/sparse_text`: a geometry-heavy sheet with almost no text.

    Grid, hatch and rectangles, but only three tiny labels (K1, K2, K3) plus
    the title block, so text-anchor matching cannot work alone. Both sides
    carry the same drawing (Rev C vs Rev D re-issue). Returns *root*.
    """
    root = Path(root)
    write_drawing_pdf(
        root / "old" / "A-101-RevC.pdf",
        drawing_no="A-101",
        revision="C",
        labels=["K1", "K2", "K3"],
        label_size=5.0,
        circles=False,
    )
    write_drawing_pdf(
        root / "new" / "A-101-RevD.pdf",
        drawing_no="A-101",
        revision="D",
        labels=["K1", "K2", "K3"],
        label_size=5.0,
        circles=False,
    )
    return root


def build_alignment_impossible(root: str | Path) -> Path:
    """`10_alignment/impossible`: two genuinely different drawings.

    The old sheet is A-101 with the RM label set, grid, hatch and circles;
    the new sheet is a different drawing: S-101, an alternative label set
    (``SEC-A1``..``SEC-A16`` on a different matrix), a solid block, a dense
    45-degree lattice instead of the grid and no circles. No alignment of the
    pair deserves to pass the quality gate. Returns *root*.
    """
    root = Path(root)
    old_dir = root / "old"
    new_dir = root / "new"

    write_drawing_pdf(old_dir / "A-101-RevC.pdf", drawing_no="A-101", revision="C")

    sec_labels = [f"SEC-A{n}" for n in range(1, 17)]
    write_drawing_pdf(
        new_dir / "S-101-RevA.pdf",
        drawing_no="S-101",
        revision="A",
        scale="1 : 100",
        labels=sec_labels,
        circles=False,
        variant=1,
    )
    return root


#: The nine `10_alignment` cases: (folder name under 10_alignment, builder).
ALIGNMENT_CASES = (
    ("clean_pair", build_alignment_clean),
    ("shifted", build_alignment_shifted),
    ("rescaled", build_alignment_rescaled),
    ("rotated", build_alignment_rotated),
    ("page_rotated", build_alignment_page_rotated),
    ("scanned", build_alignment_scanned),
    ("no_grid", build_alignment_no_grid),
    ("sparse_text", build_alignment_sparse_text),
    ("impossible", build_alignment_impossible),
)
