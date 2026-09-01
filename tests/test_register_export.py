"""The Excel register export: four sheets, readable by a person."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from engine.core.enums import IssueType, RegisterStatus, Side
from engine.core.workspace import create_workspace
from engine.register.reconciler import reconcile
from engine.report.register_xlsx import write_register, write_register_to_workspace
from tests.test_reconciler import sheet


@pytest.fixture
def register(tmp_path: Path):
    """A register with one of everything worth exporting."""
    old = [
        sheet("A-101", "C", side=Side.OLD, content="a", title="GROUND FLOOR PLAN"),
        sheet("A-102", "C", side=Side.OLD, content="same"),
        sheet("A-103", "C", side=Side.OLD, content="c"),
        sheet("A-104", "C", side=Side.OLD, content="d"),
    ]
    new = [
        sheet("A-101", "D", content="b", title="GROUND FLOOR PLAN"),
        sheet("A-102", "C", content="same"),
        sheet("A-103", "C", content="changed"),  # same rev, different file
        sheet("A-900", "A", content="e"),
        sheet(None, "A", filename="scan0001.pdf"),
    ]
    result = reconcile(old, new, issue_type=IssueType.FULL)

    quarantine = [
        {
            "filename": "locked.pdf",
            "label": "Password-protected",
            "advice": "Ask the sender for a copy without a password, then scan again.",
            "path": r"D:\new\locked.pdf",
        }
    ]
    return result, quarantine


def load(path: Path) -> openpyxl.Workbook:
    return openpyxl.load_workbook(path)


# ── Structure ──────────────────────────────────────────────────────────


def test_the_workbook_has_the_four_sheets(register, tmp_path: Path):
    result, quarantine = register
    path = write_register(
        tmp_path / "Drawing_Register.xlsx", result.rows, result.summary, quarantine=quarantine
    )

    book = load(path)
    assert book.sheetnames == ["Summary", "Register", "Needs attention", "Not readable"]


def test_the_summary_can_be_pasted_into_an_email(register, tmp_path: Path):
    result, _ = register
    path = write_register(
        tmp_path / "r.xlsx",
        result.rows,
        result.summary,
        context={
            "project_name": "Al Basateen Farm - Villa",
            "old_folder": r"D:\old",
            "new_folder": r"D:\new",
            "profile": "keo",
            "version": "0.1.0",
        },
    )

    values = [
        str(cell.value)
        for row in load(path)["Summary"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    text = " | ".join(values)

    assert "Al Basateen Farm - Villa" in text
    assert r"D:\old" in text and r"D:\new" in text
    assert "Full issue" in text
    assert "keo" in text
    assert result.summary.sentence in text


def test_the_register_sheet_lists_every_row(register, tmp_path: Path):
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    sheet_obj = load(path)["Register"]
    assert sheet_obj.max_row == len(result.rows) + 1  # a header plus the rows

    headings = [cell.value for cell in sheet_obj[1]]
    assert headings[0] == "Drawing no."
    assert "How the number was found" in headings
    assert "Note" in headings


def test_the_header_is_frozen_and_filterable(register, tmp_path: Path):
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)
    sheet_obj = load(path)["Register"]

    assert sheet_obj.freeze_panes == "A2"
    assert sheet_obj.auto_filter.ref is not None


def test_the_source_of_each_number_is_shown(register, tmp_path: Path):
    """So the reader can judge which rows to trust."""
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    sheet_obj = load(path)["Register"]
    headings = [cell.value for cell in sheet_obj[1]]
    column = headings.index("How the number was found") + 1
    values = {
        sheet_obj.cell(row=index, column=column).value for index in range(2, sheet_obj.max_row + 1)
    }

    assert "Title block" in values


# ── Needs attention ────────────────────────────────────────────────────


def test_the_attention_sheet_holds_only_the_flagged_rows(register, tmp_path: Path):
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    sheet_obj = load(path)["Needs attention"]
    flagged = [row for row in result.rows if row.needs_attention]

    assert sheet_obj.max_row == len(flagged) + 1
    assert flagged  # the fixture has some


def test_each_flagged_row_explains_what_to_do(register, tmp_path: Path):
    """ "Same revision, different file" must read as a sentence, not a code."""
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    sheet_obj = load(path)["Needs attention"]
    notes = " ".join(
        str(sheet_obj.cell(row=index, column=4).value or "")
        for index in range(2, sheet_obj.max_row + 1)
    )

    assert "without a revision change" in notes
    assert "Compare it" in notes


def test_an_empty_attention_sheet_says_so(tmp_path: Path):
    result = reconcile(
        [sheet("A-101", "C", side=Side.OLD, content="same")],
        [sheet("A-101", "C", content="same")],
        issue_type=IssueType.FULL,
    )
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    assert load(path)["Needs attention"].cell(row=2, column=1).value == "Nothing needs attention."


# ── Quarantine ─────────────────────────────────────────────────────────


def test_unreadable_files_are_listed_with_their_reason(register, tmp_path: Path):
    result, quarantine = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary, quarantine=quarantine)

    sheet_obj = load(path)["Not readable"]
    assert sheet_obj.cell(row=2, column=1).value == "locked.pdf"
    assert sheet_obj.cell(row=2, column=2).value == "Password-protected"
    assert "Ask the sender" in str(sheet_obj.cell(row=2, column=3).value)


def test_no_quarantine_says_everything_was_read(register, tmp_path: Path):
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    assert "successfully" in str(load(path)["Not readable"].cell(row=2, column=1).value)


# ── Never overwrite ────────────────────────────────────────────────────


def test_an_existing_register_is_not_overwritten(register, tmp_path: Path):
    """It may already have been sent to the design team."""
    result, _ = register
    first = write_register(tmp_path / "Drawing_Register.xlsx", result.rows, result.summary)
    second = write_register(tmp_path / "Drawing_Register.xlsx", result.rows, result.summary)

    assert first.name == "Drawing_Register.xlsx"
    assert second.name == "Drawing_Register (2).xlsx"
    assert first.exists() and second.exists()


# ── Into the workspace ─────────────────────────────────────────────────


def test_the_register_lands_in_the_register_folder(register, tmp_path: Path):
    result, quarantine = register
    workspace = create_workspace(tmp_path / "Compare_RevC_to_RevD")

    path = write_register_to_workspace(
        workspace, result.rows, result.summary, quarantine=quarantine
    )

    assert path.parent == workspace.register_dir
    assert path.name == "Drawing_Register.xlsx"
    assert path.is_file()


def test_a_status_reads_as_words_not_a_code(register, tmp_path: Path):
    result, _ = register
    path = write_register(tmp_path / "r.xlsx", result.rows, result.summary)

    sheet_obj = load(path)["Register"]
    statuses = {
        str(sheet_obj.cell(row=index, column=5).value) for index in range(2, sheet_obj.max_row + 1)
    }

    assert "Same revision, different file" in statuses
    assert str(RegisterStatus.SAME_REV_DIFFERENT_FILE) not in statuses
